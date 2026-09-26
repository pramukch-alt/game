import json
import openpyxl
import os

wb = openpyxl.load_workbook('D:/projects/Game/DrinkTalk/Source/100_คำถาม_วงเหล้า_3_Categories.xlsx')
data = {}
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    questions = []
    for r in range(2, ws.max_row + 1):
        num = ws.cell(r, 1).value
        text = ws.cell(r, 2).value
        if text:
            clean_text = str(text).strip()
            questions.append({
                'id': int(num) if num is not None else len(questions) + 1,
                'text': clean_text,
                'category': sheetname,
                'isSpecial': False
            })
    data[sheetname] = questions

special_cards = [
    {
        'id': 101,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🍺 ดื่ม 1 Drink!',
        'text': 'คนจั่วจัดไป 1 อึก/ช็อตทันที เพื่อเปิดทางความสดชื่นของค่ำคืนนี้!',
        'badge': 'ดื่มเปิดทาง',
        'icon': '🍺',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 102,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🍻 แจก 1 Drink!',
        'text': 'ชี้เพื่อนในวง 1 คนที่คุณอยากให้ดื่ม คนนั้นต้องดื่ม 1 อึกทันที!',
        'badge': 'แจกดริ๊งก์',
        'icon': '🍻',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 103,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🥂 ชนแก้วทั้งวง!',
        'text': 'ทุกคนในวงยกแก้วขึ้นมา ชนแก้วแล้วดื่มพร้อมกัน 1 อึก เพื่อมิตรภาพ!',
        'badge': 'ชนแก้วทั้งโต๊ะ',
        'icon': '🥂',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 104,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🥃 ดวล Drink เป่ายิ้งฉุบ!',
        'text': 'เลือกเพื่อนคนตรงข้ามมาเป่ายิ้งฉุบ 1 ตา ใครแพ้ต้องดื่ม 1 ช็อต!',
        'badge': 'ดวล 1-on-1',
        'icon': '🥃',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 105,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🍹 เพื่อนซ้าย-ขวาดื่ม!',
        'text': 'เพื่อนที่นั่งอยู่ทางซ้ายและทางขวาของคนจั่ว ดื่มคนละ 1 อึกทันที!',
        'badge': 'เพื่อนบ้านโดน',
        'icon': '🍹',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 106,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '🍷 ดื่มตามเพศ!',
        'text': 'คนจั่วมีสิทธิ์เลือกเพศ "ผู้ชาย" หรือ "ผู้หญิง" ทุกคนที่เป็นเพศนั้นต้องดื่ม 1 อึก!',
        'badge': 'เพศที่เลือกดื่ม',
        'icon': '🍷',
        'color': '#00F0FF',
        'isSpecial': True
    },
    {
        'id': 107,
        'category': 'Special Cards',
        'type': 'drink',
        'title': '❤️ คนมีแฟน / คนโสด ดื่ม!',
        'text': 'คนจั่วเลือกโหมด "คนมีแฟนดื่ม" หรือ "คนโสดดื่ม" ใครเข้าเกณฑ์ดื่ม 1 อึก!',
        'badge': 'สถานะหัวใจ',
        'icon': '❤️',
        'color': '#FF2E93',
        'isSpecial': True
    },
    {
        'id': 108,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '⏱️ มินิเกม: ต่อเพลงหมวดรัก!',
        'text': 'วนตามเข็มนาฬิการ้องเพลงที่มีคำว่า "รัก" คนละท่อน ใครติดเกิน 5 วินาที ดื่ม 1 ช็อต!',
        'badge': 'มินิเกมต่อเพลง',
        'icon': '🎵',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 109,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '🔢 มินิเกม: นับเลขห้าม 7!',
        'text': 'นับเลขวน 1 ถึง 30 ห้ามพูดเลขที่มี 7 หรือหารด้วย 7 ลงตัว (ให้ตบมือแทน) ใครหลุดพูดดื่ม!',
        'badge': 'มินิเกมนับเลข',
        'icon': '🔢',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 110,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '🤐 มินิเกม: ห้ามพูด ใช่/ไม่ใช่!',
        'text': 'มีผล 1 รอบวง ใครเผลอหลุดพูดคำว่า "ใช่" หรือ "ไม่ใช่" โดนดื่ม 1 อึกทันที!',
        'badge': 'มินิเกมคำต้องห้าม',
        'icon': '🤐',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 111,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '👀 มินิเกม: จ้องตาพิฆาต!',
        'text': 'เลือกเพื่อน 1 คนมานั่งจ้องตากัน 15 วินาที ใครหลุดขำหรือกะพริบตาก่อน ดื่ม 1 อึก!',
        'badge': 'มินิเกมประลองสายตา',
        'icon': '👀',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 112,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '🤫 มินิเกม: 2 เรื่องจริง 1 เรื่องโกหก!',
        'text': 'คนจั่วเล่าเรื่องเกี่ยวกับตัวเอง 3 เรื่อง (จริง 2 โกหก 1) ให้คนในวงทาย ถ้าเพื่อนทายถูกคนเล่าดื่ม ถ้าทายผิดคนทายดื่ม!',
        'badge': 'มินิเกมจับโกหก',
        'icon': '🤫',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 113,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '☝️ มินิเกม: Most Likely To!',
        'text': 'คนจั่วนับ 1 2 3 แล้วให้ทุกคนชี้คนที่ "น่าจะเมาปลิ้นตื่นมาจำอะไรไม่ได้ที่สุด" คนโดนชี้มากสุดดื่ม!',
        'badge': 'มินิเกมชี้เป้า',
        'icon': '☝️',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 114,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '📱 มินิเกม: เปิดแกลเลอรี่ล่าสุด!',
        'text': 'เปิดรูปถ่ายล่าสุดในเครื่องให้เพื่อนดู หรือถ้าไม่กล้าเปิดเลือกดื่ม 2 ดริ๊งก์!',
        'badge': 'วัดใจเปิดรูป',
        'icon': '📱',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 115,
        'category': 'Special Cards',
        'type': 'minigame',
        'title': '🎭 มินิเกม: ใบ้คำด้วยท่าทาง!',
        'text': 'คนทางขวาบอกคำลับ 1 คำให้คนจั่วใบ้ด้วยท่าทาง ห้ามออกเสียง ถ้าเพื่อนทายถูกใน 30 วิ คนบอกดื่ม ถ้าทายไม่ถูกคนใบ้ดื่ม!',
        'badge': 'มินิเกมใบ้คำ',
        'icon': '🎭',
        'color': '#FFB800',
        'isSpecial': True
    },
    {
        'id': 116,
        'category': 'Special Cards',
        'type': 'rule',
        'title': '🔄 สลับทิศทางวง!',
        'text': 'ทิศทางการเล่นและคนจั่วใบถัดไปเปลี่ยนเป็นทวนเข็มนาฬิกาทันที!',
        'badge': 'ทิศทางย้อนกลับ',
        'icon': '🔄',
        'color': '#7928CA',
        'isSpecial': True
    },
    {
        'id': 117,
        'category': 'Special Cards',
        'type': 'rule',
        'title': '🎯 สลับตัวประกัน!',
        'text': 'คนจั่วได้รับการยกเว้น และสามารถเลือกชี้เพื่อน 1 คนให้รับคำถามใบถัดไปแทนได้เลย!',
        'badge': 'โยนการ์ดให้เพื่อน',
        'icon': '🎯',
        'color': '#7928CA',
        'isSpecial': True
    },
    {
        'id': 118,
        'category': 'Special Cards',
        'type': 'rule',
        'title': '🎲 วงล้อชะตากรรม!',
        'text': 'กดปุ่มสุ่มคนในวงทันที! คนที่โดนสุ่มต้องดื่ม 1 อึกแล้วเป็นคนตอบคำถามถัดไป!',
        'badge': 'สุ่มผู้โชคดี',
        'icon': '🎲',
        'color': '#7928CA',
        'isSpecial': True
    },
    {
        'id': 119,
        'category': 'Special Cards',
        'type': 'rule',
        'title': '👑 ราชาสั่งการ (King)!',
        'text': 'คนจั่วได้สิทธิ์เป็น King สั่งคำสั่งขำๆ ให้คนในวงทำ 1 อย่าง (เช่น ให้ยืนเต้น 5 วิ หรือพูดลงท้ายด้วย "เจ้าค่ะ")!',
        'badge': 'คำสั่งราชา',
        'icon': '👑',
        'color': '#7928CA',
        'isSpecial': True
    },
    {
        'id': 120,
        'category': 'Special Cards',
        'type': 'rule',
        'title': '🛡️ การ์ดรอดชีวิต (Safe Card)!',
        'text': 'เก็บการ์ดใบนี้ไว้กับตัว สามารถใช้ข้ามคำถามที่ไม่อยากตอบ หรือยกเลิกการดื่มได้ 1 ครั้ง!',
        'badge': 'การ์ดป้องกัน',
        'icon': '🛡️',
        'color': '#00F0FF',
        'isSpecial': True
    }
]

data['Special Cards'] = special_cards

os.makedirs('D:/projects/Game/DrinkTalk/data', exist_ok=True)
with open('D:/projects/Game/DrinkTalk/data/questions.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

with open('D:/projects/Game/DrinkTalk/data/questions.js', 'w', encoding='utf-8') as f:
    f.write('window.DRINKTALK_QUESTIONS = ' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n')

print('Successfully exported all categories and special cards!')
